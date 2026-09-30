"""Build the Styllano color-system workbook from Shopify's public products.json dump."""
import json, re, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else 'all.json'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'styllano_colors.xlsx'
COLOR_OPTS = {'Color', 'Colors', 'color', 'ColorB', 'الالوان'}
ADMIN = 'https://admin.shopify.com/store/styllano/products/'

# raw (lowercase) -> (standard name, reason). reason: ok | spelling | arabic | synonym | review
S = {}
def add(std, reason, *raws):
    for r in raws: S[r.lower()] = (std, reason)

add('Black', 'ok', 'black'); add('Black', 'arabic', 'أسود', 'اسود')
add('Black', 'review', 'matte black', 'black burnished', 'black suede')
add('White', 'ok', 'white'); add('White', 'spelling', 'wihte'); add('White', 'arabic', 'ابيض')
add('Off White', 'ok', 'off white'); add('Off White', 'spelling', 'off-white')
add('Gray', 'ok', 'gray'); add('Gray', 'spelling', 'grey'); add('Gray', 'arabic', 'رمادي'); add('Gray', 'review', 'رصاصي')
add('Dark Gray', 'ok', 'dark gray'); add('Dark Gray', 'spelling', 'dark grey')
add('Light Gray', 'ok', 'light gray'); add('Light Gray', 'spelling', 'light grey')
add('Beige', 'ok', 'beige'); add('Beige', 'spelling', 'baige', 'biege', 'bagie'); add('Beige', 'arabic', 'ييج', 'بيج')
add('Brown', 'ok', 'brown'); add('Brown', 'spelling', 'brawn', 'brwon'); add('Brown', 'arabic', 'بني'); add('Brown', 'review', 'brown suede')
add('Burgundy', 'ok', 'burgundy'); add('Burgundy', 'spelling', 'burgandy', 'burgany')
add('Navy', 'ok', 'navy'); add('Navy', 'synonym', 'navy blue'); add('Navy', 'arabic', 'كحلي')
add('Blue', 'ok', 'blue'); add('Dark Blue', 'ok', 'dark blue'); add('Light Blue', 'ok', 'light blue')
add('Baby Blue', 'ok', 'baby blue'); add('Dark Baby Blue', 'ok', 'dark baby blue')
add('Denim Blue', 'ok', 'denim blue'); add('Denim Blue', 'synonym', 'blue denim', 'denim'); add('Denim Blue', 'review', 'jeans')
add('Light Denim', 'synonym', 'light blue denim'); add('Dark Denim', 'synonym', 'dark blue denim'); add('Black Denim', 'ok', 'black denim')
add('Havana', 'ok', 'havana'); add('Havana', 'spelling', 'havan'); add('Havana', 'arabic', 'هافان'); add('Havana', 'review', 'havana brown')
add('Cashmere', 'ok', 'cashmere'); add('Cashmere', 'spelling', 'cashmir', 'kashmir')
add('Camel', 'ok', 'camel'); add('Camel', 'spelling', 'gamal', 'gamel')
add('Genzari', 'review', 'genzari', 'janazary')
add('Coffee', 'ok', 'coffee'); add('Coffee', 'spelling', 'caffe', 'coffe'); add('Cocoa', 'review', 'cocoa')
add('Mauve', 'ok', 'mauve'); add('Mauve', 'spelling', 'mouve', 'muve')
add('Mustard', 'ok', 'mustard'); add('Mustard', 'spelling', 'mustarda')
add('Mint Green', 'ok', 'mint green'); add('Mint Green', 'review', 'mint')
add('Kiwi', 'ok', 'kiwi'); add('Kiwi', 'synonym', 'kiwi green')
add('Olive', 'ok', 'olive'); add('Olive', 'review', 'olive green')
add('Dark Olive', 'ok', 'dark olive'); add('Light Olive', 'ok', 'light olive')
add('Petrol', 'ok', 'petrol'); add('Petrol', 'synonym', 'petrol blue')
add('Pepsi', 'ok', 'pepsi'); add('Pepsi', 'synonym', 'pepsi blue')
add('Turquoise', 'ok', 'turquoise'); add('Turquoise', 'synonym', 'turquoise blue')
add('Pink', 'ok', 'pink'); add('Light Pink', 'ok', 'light pink'); add('Rose', 'review', 'rose'); add('Dusty Rose', 'review', 'dusty rose')
add('Yellow', 'ok', 'yellow'); add('Neon Yellow', 'ok', 'neon yellow'); add('Lemon Yellow', 'review', 'lemon yellow')
add('Khaki', 'ok', 'khaki'); add('Silver', 'ok', 'silver'); add('Silver', 'spelling', 'selver')
add('Gold', 'ok', 'gold'); add('Gold', 'review', 'gold 18', 'gold 20', 'cold'); add('Rose Gold', 'ok', 'rose gold')
add('Mango', 'ok', 'mango'); add('Mango', 'arabic', 'منجاوي')
add('Salmon', 'ok', 'salmon'); add('Salmon', 'review', 'simon')
add('Cumin', 'review', 'كموني')
for c in ['Aqua', 'Brick', 'Dark Green', 'Fuchsia', 'Green', 'Orange', 'Peach', 'Pistachio', 'Purple',
          'Red', 'Taupe', 'Teal', 'Terracotta', 'Camouflage Green', 'Sand']:
    add(c, 'ok', c.lower())
add('Sand', 'ok', 'sand'); add('Desert Beige', 'review', 'desert beige')
add('Firani', 'review', 'firani'); add('Shania', 'review', 'shania'); add('Burlap', 'review', 'burlap')
add('Watermelon', 'review', 'watermelon')

REASON_AR = {'ok': 'سليم', 'spelling': 'غلط إملائي/كتابة', 'arabic': 'عربي ← إنجليزي',
             'synonym': 'اسم تاني لنفس اللون', 'review': 'محتاج تبص عليه بعينك',
             'combo': 'لون مزدوج — توحيد الصيغة', 'notcolor': 'مش لون — انقله لـoption تاني'}

NOT_COLOR = re.compile(r'^(pack|bundle|a[a-d]$|cat$|tiger$|stars$)', re.I)
SPLIT = re.compile(r'\s+with\s+|\s+and\s+|\s*&\s*|\s*×\s*|\s+x\s+|\s*\*\s*|\s*/\s*|\s+–\s+|\s+-\s+|(?<=[؀-ۿ])X(?=[؀-ۿ])', re.I)
DETAIL = re.compile(r'\b(logo|stripes|striped|sole|accents?|mark|dual-face|leather accent|stars|\d+)\b', re.I)

def norm(v): return re.sub(r'\s+', ' ', v).strip().rstrip('.').strip()

def standardize(raw):
    """-> (standard, reason, note)"""
    v = norm(raw)
    if NOT_COLOR.match(v): return ('— (مش لون)', 'notcolor', 'Pack / تصميم / باترن — مكانه option منفصل زي "Pack" أو "Design"')
    if v.lower() in S:
        std, why = S[v.lower()]
        note = ''
        if v.lower() in ('gold 18', 'gold 20'): note = 'الرقم ده عيار مش لون — حطه في option تاني'
        elif v.lower() == 'cold': note = 'غالباً غلطة في Gold — اتأكد'
        elif 'suede' in v.lower() or v.lower() in ('matte black', 'black burnished'): note = 'دي خامة/تشطيب مش لون'
        elif v == std and why == 'ok': why = 'ok'
        elif v != std and why == 'ok': why = 'spelling'  # case differences
        return (std, why, note)
    parts = [p for p in SPLIT.split(v) if p and p.strip()]
    details = DETAIL.findall(v)
    out = []
    for p in parts:
        q = norm(DETAIL.sub('', p))
        if not q: continue
        if q.lower() in S: out.append(S[q.lower()][0])
        elif q.lower() in ('leather', 'tiger'): details.append(q)
        else: out.append(q.title())
    out = list(dict.fromkeys(out))
    note = ('التفاصيل (' + ', '.join(sorted(set(d.title() for d in details))) + ') مكانها الوصف مش اسم اللون') if details else ''
    if len(out) == 1:
        return (out[0], 'review' if details else 'synonym', note)
    if len(out) >= 2:
        return (' & '.join(out), 'combo', note)
    return (v.title(), 'review', 'مش متعرف')

prods = json.load(open(SRC))
rows = []  # per product x color value
opt_names = collections.Counter()
for p in prods:
    for o in p['options']:
        opt_names[o['name']] += 1
        if o['name'] not in COLOR_OPTS: continue
        key = 'option%d' % o['position']
        for val in o['values']:
            vs = [x for x in p['variants'] if x.get(key) == val]
            skus = list(dict.fromkeys((x.get('sku') or '').strip() for x in vs if (x.get('sku') or '').strip()))
            urls = [s for s in skus if s.startswith('http')]
            texts = [s for s in skus if not s.startswith('http')]
            vend_color = ''
            if texts:
                m = re.search(r' - (.+?)\s*\d*$', texts[0])
                vend_color = m.group(1).strip() if m else ''
            std, why, note = standardize(val)
            rows.append(dict(title=p['title'], vendor=p['vendor'], admin=ADMIN + str(p['id']), opt=o['name'],
                             raw=val, std=std, why=why, note=note,
                             vlink=' , '.join(urls[:3]) or (texts[0] if texts else ''), vcolor=vend_color))

# ---------- workbook ----------
F = 'Arial'
hdr_font = Font(name=F, bold=True, color='FFFFFF'); hdr_fill = PatternFill('solid', fgColor='2F3E4E')
body = Font(name=F, size=10); link = Font(name=F, size=10, color='0563C1', underline='single')
fill_in = PatternFill('solid', fgColor='FFFF00')
why_fill = {'ok': 'E2EFDA', 'spelling': 'FCE4D6', 'arabic': 'FCE4D6', 'synonym': 'FFF2CC',
            'review': 'F8CBAD', 'combo': 'DDEBF7', 'notcolor': 'D9D9D9'}

def sheet(ws, headers, widths):
    ws.sheet_view.rightToLeft = False
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        c = ws.cell(1, i); c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

wb = Workbook()

# Sheet 1: standard colors
ws1 = wb.active; ws1.title = 'Standard Colors'
std_raws = collections.defaultdict(collections.Counter); std_why = collections.defaultdict(set); std_notes = collections.defaultdict(set)
for r in rows:
    std_raws[r['std']][r['raw']] += 1; std_why[r['std']].add(r['why'])
    if r['note']: std_notes[r['std']].add(r['note'])
def kind(s):
    w = std_why[s]
    if 'notcolor' in w: return 3
    if 'combo' in w: return 2
    return 1
order = sorted(std_raws, key=lambda s: (kind(s), -sum(std_raws[s].values()), s))
sheet(ws1, ['Standard name (الاسم الموحد)', 'النوع', 'عدد المنتجات', 'عدد الأسماء المختلفة', 'الأسماء الموجودة حالياً (وعدد المنتجات)', 'ملاحظات / محتاج مراجعة'],
      [28, 14, 12, 12, 70, 60])
n = len(rows) + 1
for s in order:
    names = std_raws[s]
    typ = {1: 'لون', 2: 'لون مزدوج', 3: 'مش لون'}[kind(s)]
    notes = sorted(std_notes[s])
    if 'review' in std_why[s]: notes.insert(0, '⚠ فيه أسماء محتاجة تتأكد منها بعينك')
    ws1.append([s, typ, None, len(names), '  |  '.join(f'{k} ({v})' for k, v in names.most_common()), ' / '.join(notes)])
    r = ws1.max_row
    ws1.cell(r, 3).value = f"=COUNTIF(Products!$F$2:$F${n},A{r})"
    for c in range(1, 7): ws1.cell(r, c).font = body; ws1.cell(r, c).alignment = Alignment(wrap_text=True, vertical='top')
    ws1.cell(r, 1).font = Font(name=F, size=10, bold=True)
    if len(names) > 1 or 'review' in std_why[s]:
        ws1.cell(r, 4).fill = PatternFill('solid', fgColor='FCE4D6')
ws1.auto_filter.ref = f'A1:F{ws1.max_row}'

# Sheet 2: every raw value
ws2 = wb.create_sheet('All Values')
sheet(ws2, ['القيمة زي ما هي في شوبيفاي', 'عدد المنتجات', 'الاسم الموحد المقترح', 'الحالة', 'ملاحظة'], [42, 12, 28, 26, 60])
raw_info = {}
for r in rows: raw_info[r['raw']] = (r['std'], r['why'], r['note'])
for raw in sorted(raw_info, key=lambda x: (raw_info[x][0].lower(), x.lower())):
    std, why, note = raw_info[raw]
    ws2.append([raw, None, std, REASON_AR[why], note]); r = ws2.max_row
    ws2.cell(r, 2).value = f'=COUNTIF(Products!$E$2:$E${n},A{r})'
    for c in range(1, 6): ws2.cell(r, c).font = body
    ws2.cell(r, 4).fill = PatternFill('solid', fgColor=why_fill[why])
ws2.auto_filter.ref = f'A1:E{ws2.max_row}'

# Sheet 3: products
ws3 = wb.create_sheet('Products')
sheet(ws3, ['المنتج', 'Vendor (شوبيفاي)', 'لينك الأدمن', 'اسم الـOption', 'اللون في شوبيفاي', 'الاسم الموحد المقترح', 'الحالة',
            'لينك/SKU Vendoor', 'اللون في Vendoor (يتملى)', 'ملاحظة'], [40, 16, 14, 11, 26, 24, 22, 38, 22, 45])
for x in sorted(rows, key=lambda x: (x['title'].lower(), x['raw'])):
    ws3.append([x['title'], x['vendor'], 'فتح', x['opt'], x['raw'], x['std'], REASON_AR[x['why']], x['vlink'], x['vcolor'] or None, x['note']])
    r = ws3.max_row
    for c in range(1, 11): ws3.cell(r, c).font = body
    ws3.cell(r, 3).hyperlink = x['admin']; ws3.cell(r, 3).font = link
    if x['vlink'].startswith('http'):
        ws3.cell(r, 8).hyperlink = x['vlink'].split(' , ')[0]; ws3.cell(r, 8).font = link
    if x['opt'] != 'Color': ws3.cell(r, 4).fill = PatternFill('solid', fgColor='F8CBAD')
    ws3.cell(r, 7).fill = PatternFill('solid', fgColor=why_fill[x['why']])
    if not x['vcolor']: ws3.cell(r, 9).fill = fill_in
ws3.auto_filter.ref = f'A1:J{ws3.max_row}'

# Sheet 4: option names
ws4 = wb.create_sheet('Option Names')
sheet(ws4, ['اسم الـOption', 'عدد المنتجات', 'المقترح'], [22, 14, 50])
for k, v in opt_names.most_common():
    sug = 'Color' if k in COLOR_OPTS else ('Size' if k in ('Size', 'size', 'المقاس', 'Shoe size') else '')
    if k == 'Shoe size': sug = 'Size (أو سيبه لو قاصد تفرق مقاسات الجزم)'
    if k == 'Title': sug = 'منتج من غير variants — سيبه'
    ws4.append([k, v, ('✓ سليم' if k == sug else sug)])
    for c in range(1, 4): ws4.cell(ws4.max_row, c).font = body
ws4.append([]); ws4.append(['شوبيفاي بيعمل فلتر منفصل لكل اسم option، فـColor و Colors بيطلعوا للزبون كفلترين. وحّدهم كلهم لـColor.'])
ws4.cell(ws4.max_row, 1).font = Font(name=F, size=10, italic=True)

# Legend sheet
ws0 = wb.create_sheet('اقرأني', 0)
lines = [
    ('Styllano — Color System', True),
    (f'اتسحب من شوبيفاي (products.json): {len(prods)} منتج، {len(rows)} صف لون، {len(raw_info)} قيمة لون مختلفة ← {len([s for s in order if kind(s)==1])} لون موحد + {len([s for s in order if kind(s)==2])} لون مزدوج.', False),
    ('', False),
    ('Standard Colors: كل لون موحد مرة واحدة، وجنبه كل الأسماء اللي متكتب بيها حالياً. ده الجدول الأساسي.', False),
    ('All Values: كل قيمة لون زي ما هي في شوبيفاي ← الاسم الموحد المقترح.', False),
    ('Products: كل منتج × لون، مع لينك الأدمن ولينك Vendoor. عمود "اللون في Vendoor" الأصفر هيتملى لما ملف Vendoor يوصل.', False),
    ('Option Names: أسماء الـoption نفسها (Color / Colors / الالوان...) — لازم تتوحد.', False),
    ('', False),
    ('ألوان عمود الحالة:', True),
    ('أخضر = سليم | برتقالي فاتح = غلط كتابة أو عربي | أصفر = اسم تاني لنفس اللون | أحمر فاتح = محتاج تبص عليه بعينك | أزرق = لون مزدوج | رمادي = مش لون', False),
    ('', False),
    ('قاعدة الألوان المزدوجة المقترحة: "Black & White" (اللون الأساسي الأول). مش "/" عشان شوبيفاي بيستخدم "/" بين الـoptions في اسم الـvariant.', False),
    ('تفاصيل زي Logo / Stripes / Sole مكانها الوصف أو الصور، مش اسم اللون.', False),
    ('الأسماء في "محتاج مراجعة" اقتراحات مني — لو الدرجتين مختلفين فعلاً قدام الزبون، ما تدمجهمش.', False),
]
ws0.column_dimensions['A'].width = 130
for t, b in lines:
    ws0.append([t]); c = ws0.cell(ws0.max_row, 1)
    c.font = Font(name=F, size=14 if t.startswith('Styllano') else 10, bold=b); c.alignment = Alignment(wrap_text=True)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('rows', len(rows), 'raw', len(raw_info), 'std', len(order))
