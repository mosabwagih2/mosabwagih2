"""Add Vendoor-vs-Shopify color matching to the color-system workbook.

Usage: python merge_vendoor.py <workbook with Vendoor column filled> <output.xlsx>
Reads Products!I (Vendoor color, Arabic, as copied from Vendoor), maps it to the
standard English names, and compares with Products!F (Shopify standard name).
"""
import re, sys, collections
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
F = 'Arial'
body = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
hdr_font = Font(name=F, bold=True, color='FFFFFF'); hdr_fill = PatternFill('solid', fgColor='2F3E4E')
fill = lambda c: PatternFill('solid', fgColor=c)

# Vendoor Arabic term (normalized) -> (standard English, needs_review, note)
V = {}
def v(std, *terms, review=False, note=''):
    for t in terms: V[t] = (std, review, note)
v('Black', 'اسود', 'سوداء', 'اسود ناعم'); v('Black', 'اسود فرنيه', review=True, note='فرنيه = تشطيب لامع، مش لون')
v('White', 'ابيض', 'بيضاء'); v('Off White', 'اوف وايت')
v('Gray', 'رمادي'); v('Gray', 'رصاصي', review=True, note='Vendoor بيستخدم رمادي ورصاصي الاتنين — 21 من 23 في شوبيفاي Gray/Grey و 2 Dark Gray')
v('Gray', 'حديدي', review=True, note='حديدي غالباً أغمق — ممكن Dark Gray')
v('Dark Gray', 'رمادي غامق', 'رصاصي غامق'); v('Light Gray', 'رمادي فاتح', 'رصاصي فاتح')
v('Heather Gray', 'رمادي شانيه', review=True, note='شانيه = chiné/melange — في شوبيفاي مكتوب Shania')
v('Navy', 'كحلي'); v('Beige', 'بيج'); v('Brown', 'بني'); v('Blue', 'ازرق'); v('Havana', 'هافان')
v('Olive', 'زيتي', 'زتوني'); v('Dark Olive', 'زيتي غامق'); v('Light Olive', 'زيتي فاتح')
v('Red', 'احمر', 'حمراء'); v('Gold', 'جولد', 'ذهبي'); v('Silver', 'سيلفر', 'فضي', 'سلفر')
v('Pink', 'بينك'); v('Pink', 'زهري', review=True, note='زهري = وردي/بينك')
v('Rose', 'روز'); v('Burgundy', 'نبيتي', 'برجندي')
v('Baby Blue', 'لبني', 'بيبي بلو', review=True, note='لبني = Baby Blue (15) بس اتكتب Light Blue (7) — لازم اسم واحد')
v('Light Blue', 'لبني فاتح', review=True, note='اتكتب Baby Blue 3 مرات و Light Blue مرة')
v('Dark Baby Blue', 'لبني غامق')
v('Cashmere', 'كشمير'); v('Yellow', 'اصفر'); v('Camel', 'جملي'); v('Petrol', 'بترولي')
v('Mint Green', 'مينت جرين', 'منت جرين', 'منت'); v('Green', 'اخضر', 'خضراء'); v('Mauve', 'موف')
v('Turquoise', 'تركواز'); v('Turquoise', 'فيروزي', review=True, note='اتكتب Aqua و Teal و Turquoise')
v('Coffee', 'كافيه'); v('Fuchsia', 'فوشيا'); v('Pepsi', 'بيبسي'); v('Mango', 'منجاوي'); v('Cocoa', 'كاكاو')
v('Cumin', 'كموني', review=True, note='اتكتب Camel 3 مرات و Khaki مرتين — لون مختلف عنهم')
v('Genzari', 'جنزاري', 'اخضر جنزاري', review=True, note='اتكتب Genzari و Green و Dark Green و Teal و Petrol')
v('Mustard', 'مسترده'); v('Orange', 'اورانج'); v('Pistachio', 'بستشيو'); v('Purple', 'بنفسجي')
v('Desert Beige', 'صحراوي', review=True, note='اتكتب Sand و Desert Beige'); v('Sand', 'رملي')
v('Denim Blue', 'جينز', 'لون جينز'); v('Black Denim', 'جينز اسود'); v('Dark Denim', 'جينز كحلي'); v('Light Denim', 'جينز لبني')
v('Firani', 'فيراني', review=True); v('Salmon', 'سيمون'); v('Brick', 'طوبي', review=True, note='اتكتب Brick و Terracotta')
v('Kiwi', 'كيوي'); v('Burlap', 'خيش'); v('Neon Yellow', 'فسفوري'); v('Lemon Yellow', 'لموني'); v('Watermelon', 'بطيخي')
NOT_COLOR = re.compile(r'^(باندل|AA|AB|AC|AD|Cat$|مشجر$)', re.I)

def ar(t):
    for a, b in [('أ', 'ا'), ('إ', 'ا'), ('آ', 'ا'), ('ٍ', ''), ('ى', 'ي'), ('چ', 'ج'), ('ة', 'ه'), ('أببض', 'ابيض'), ('اببض', 'ابيض')]:
        t = t.replace(a, b)
    t = re.sub(r'[()]', '', t)
    t = re.sub(r'size-|[٠-٩\d]+\s*سنوات|\bone size\b|مقاس.*$', '', t, flags=re.I)
    for _ in range(2):
        t = re.sub(r'\s+(\d?X{1,3}L|X{0,3}[SML]|2X|\d{1,3})\s*$', '', t, flags=re.I)
    return re.sub(r'\s+', ' ', t).strip()

SEP = re.compile(r'\s*(?:\bتطعيم\b|\bمطعم\b|\bخطوط\b|\bوعلامه\b|\bعلامه\b|\*|×|(?<=\S)X(?=\S)|\sx\s|\sX\s|\sفي\s|\d?شريط)\s*')
DROP = {'جلد'}

def to_std(term):
    """normalized Arabic term -> (std or None, review, note)"""
    if NOT_COLOR.match(term): return ('— (مش لون)', False, '')
    if term in V: return V[term]
    parts = [p for p in SEP.split(term.replace('ستارز', '').strip()) if p and p not in DROP]
    if len(parts) == 1 and parts[0] in V:
        return (V[parts[0]][0], True, 'التفاصيل (جلد/ستارز) مش لون')
    if len(parts) > 1:
        out, rev, notes = [], False, []
        for p in parts:
            s = V.get(p)
            if not s: return (None, True, f'مش عارف "{p}"')
            if s[0] not in out: out.append(s[0])
            rev |= s[1]
        return (' & '.join(out), rev, '')
    return (None, True, 'مش متعرف')

wb = load_workbook(SRC)
P = wb['Products']
last = P.max_row
# new columns K, L, M
for col, (h, w) in enumerate([('اللون في Vendoor (موحد)', 24), ('المطابقة', 26), ('ملاحظة المطابقة', 50)], start=11):
    c = P.cell(1, col, h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical='center')
    P.column_dimensions[get_column_letter(col)].width = w

ST = {'match': ('✅ متطابق', 'C6EFCE'), 'order': ('↔ نفس الألوان، ترتيب مختلف', 'FFF2CC'),
      'diff': ('❌ مختلف', 'FFC7CE'), 'partial': ('◐ Vendoor كاتب جزء من اللون', 'FFEB9C'), 'gone': ('⛔ المنتج اتشال من Vendoor', 'FF9999'),
      'missing': ('⛔ اللون مش موجود على Vendoor', 'FF9999'), 'nocolor': ('ℹ Vendoor مش كاتب لون', 'D9D9D9'),
      'unknown': ('❓ مش متعرف', 'F8CBAD'), 'empty': ('— مفيش لينك Vendoor', 'D9D9D9'),
      'notcolor': ('ℹ مش لون (باندل/موديل)', 'D9D9D9')}

dict_use = collections.defaultdict(lambda: {'n': 0, 'raw': collections.Counter(), 'shop': collections.Counter()})
stats = collections.Counter()
for r in range(2, last + 1):
    raw_v = P.cell(r, 9).value; shop = P.cell(r, 6).value; shop_raw = P.cell(r, 5).value
    std, key, note = None, None, ''
    if not raw_v: key = 'empty'
    elif raw_v.startswith('—'):
        key = 'gone' if '404' in raw_v else 'missing' if 'مش موجود' in raw_v else 'nocolor'
    else:
        stds, rev, notes = [], False, []
        for t in raw_v.split(' / '):
            n = ar(t)
            s, rv, nt = to_std(n)
            dict_use[n]['n'] += 1; dict_use[n]['raw'][t.strip()] += 1; dict_use[n]['shop'][shop_raw] += 1
            if s is None: notes.append(nt); continue
            if s not in stds: stds.append(s)
            if nt: notes.append(nt)
        std = ' , '.join(stds) if stds else None
        if not stds: key = 'unknown'
        elif stds == ['— (مش لون)']: key = 'notcolor' if shop == '— (مش لون)' else 'diff'
        elif len(stds) == 1 and stds[0] == shop: key = 'match'
        elif len(stds) == 1 and set(stds[0].split(' & ')) == set(shop.split(' & ')): key = 'order'
        elif len(stds) == 1 and ' & ' in shop and set(stds[0].split(' & ')) < set(shop.split(' & ')): key = 'partial'
        else: key = 'diff'
        note = ' / '.join(dict.fromkeys(notes))
        if key in ('diff', 'partial'):
            note = (f'شوبيفاي: {shop_raw} ← Vendoor: {raw_v}' + (' — ' + note if note else ''))
    stats[key] += 1
    label, color = ST[key]
    P.cell(r, 11, std).font = body
    c = P.cell(r, 12, label); c.font = body; c.fill = fill(color)
    P.cell(r, 13, note or None).font = body
P.auto_filter.ref = f'A1:M{last}'

# ---- Vendoor dictionary sheet ----
if 'Vendoor Dictionary' in wb.sheetnames: del wb['Vendoor Dictionary']
D = wb.create_sheet('Vendoor Dictionary', 2)
heads = [('اسم اللون في Vendoor', 26), ('الاسم الموحد (English)', 24), ('عدد الصفوف', 11),
         ('إزاي Vendoor كاتبه بالظبط', 40), ('إزاي اتكتب في شوبيفاي', 50), ('متسق؟', 16), ('ملاحظة', 55)]
for i, (h, w) in enumerate(heads, 1):
    c = D.cell(1, i, h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True)
    D.column_dimensions[get_column_letter(i)].width = w
D.freeze_panes = 'A2'
for term, u in sorted(dict_use.items(), key=lambda kv: -kv[1]['n']):
    s, rev, nt = to_std(term)
    shop_names = u['shop']
    consistent = len(shop_names) == 1
    D.append([term, s or '❓', None, '  |  '.join(f'{k} ({n})' for k, n in u['raw'].most_common()),
              '  |  '.join(f'{k} ({n})' for k, n in shop_names.most_common()),
              '✅ اسم واحد' if consistent else f'⚠ {len(shop_names)} أسماء', nt or None])
    r = D.max_row
    D.cell(r, 3).value = u['n']
    for ci in range(1, 8): D.cell(r, ci).font = body; D.cell(r, ci).alignment = Alignment(wrap_text=True, vertical='top')
    D.cell(r, 1).font = bold
    D.cell(r, 6).fill = fill('C6EFCE' if consistent else 'FFEB9C')
    if rev: D.cell(r, 2).fill = fill('F8CBAD')
D.auto_filter.ref = f'A1:G{D.max_row}'

# ---- Problems sheet ----
if 'To Fix' in wb.sheetnames: del wb['To Fix']
T = wb.create_sheet('To Fix', 1)
cols = [(1, 'المنتج', 38), (3, 'لينك الأدمن', 12), (5, 'اللون في شوبيفاي', 22), (9, 'اللون في Vendoor', 28),
        (11, 'Vendoor (موحد)', 22), (12, 'المشكلة', 28), (13, 'التفاصيل', 55), (8, 'لينك Vendoor', 36)]
for i, (_, h, w) in enumerate(cols, 1):
    c = T.cell(1, i, h); c.font = hdr_font; c.fill = hdr_fill
    T.column_dimensions[get_column_letter(i)].width = w
T.freeze_panes = 'A2'
prio = {'gone': 0, 'missing': 1, 'diff': 2, 'unknown': 3, 'partial': 4, 'order': 5}
rev_lbl = {ST[k][0]: k for k in ST}
probs = [r for r in range(2, last + 1) if rev_lbl[P.cell(r, 12).value] in prio]
probs.sort(key=lambda r: (prio[rev_lbl[P.cell(r, 12).value]], P.cell(r, 1).value))
for r in probs:
    T.append([P.cell(r, c).value for c, _, _ in cols]); tr = T.max_row
    for i, (c, _, _) in enumerate(cols, 1):
        T.cell(tr, i).font = body
        src = P.cell(r, c)
        if src.hyperlink: T.cell(tr, i).hyperlink = src.hyperlink.target; T.cell(tr, i).font = Font(name=F, size=10, color='0563C1', underline='single')
    T.cell(tr, 6).fill = PatternFill("solid", fgColor=P.cell(r, 12).fill.fgColor.rgb)
T.auto_filter.ref = f'A1:H{T.max_row}'

# ---- Standard Colors: add Vendoor names column ----
S = wb['Standard Colors']
S.cell(1, 7, 'الاسم في Vendoor (وعدد الصفوف)').font = hdr_font; S.cell(1, 7).fill = hdr_fill
S.column_dimensions['G'].width = 45
by_std = collections.defaultdict(collections.Counter)
for term, u in dict_use.items():
    s = to_std(term)[0]
    if s: by_std[s][term] += u['n']
for r in range(2, S.max_row + 1):
    names = by_std.get(S.cell(r, 1).value)
    c = S.cell(r, 7, '  |  '.join(f'{k} ({n})' for k, n in names.most_common()) if names else None)
    c.font = body; c.alignment = Alignment(wrap_text=True, vertical='top')

# ---- readme summary ----
R = wb['اقرأني']
R.append([]); R.append(['مطابقة Vendoor']); R.cell(R.max_row, 1).font = Font(name=F, size=12, bold=True)
for k in ['match', 'order', 'partial', 'diff', 'gone', 'missing', 'unknown', 'nocolor', 'notcolor', 'empty']:
    R.append([f'{ST[k][0]}: {stats[k]} صف']); R.cell(R.max_row, 1).font = body
for t in ['To Fix: كل الصفوف اللي محتاجة تتصلح، مترتبة بالأولوية (منتجات اتشالت من Vendoor الأول).',
          'Vendoor Dictionary: كل اسم لون عربي في Vendoor ← الاسم الإنجليزي الموحد، وإزاي اتكتب في شوبيفاي. "⚠" = نفس اللون العربي اتكتب بأكتر من اسم في شوبيفاي.',
          'Products: اتضاف 3 أعمدة: اللون في Vendoor (موحد)، المطابقة، ملاحظة المطابقة.']:
    R.append([t]); R.cell(R.max_row, 1).font = body; R.cell(R.max_row, 1).alignment = Alignment(wrap_text=True)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(dict(stats))
